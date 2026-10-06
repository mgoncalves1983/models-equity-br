#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Recalcula um .xlsx no LibreOffice (headless) e varre erros de fórmula.

O openpyxl grava fórmulas sem valor calculado; sem recálculo, nenhum check com
data_only=True enxerga número. Este script abre o arquivo num LibreOffice com perfil
temporário configurado para SEMPRE recalcular na carga, salva de volta em .xlsx e
conta #REF!, #NAME?, #DIV/0!, #VALUE!, #N/A, #NUM!, #NULL!.

Uso:
  python3 recalc.py modelo.xlsx              # recalcula e grava no próprio arquivo
  python3 recalc.py modelo.xlsx --no-save    # recalcula numa cópia temporária e só reporta
Saída: JSON com total_errors, contagem por tipo e até 20 células de exemplo por tipo.
Código de saída: 0 sem erros; 1 com erros de fórmula; 2 falha ao recalcular.
ATENÇÃO: nunca recalcular arquivo com histórico de corrupção (abrir no Excel antes).
"""
import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from config import find_soffice  # noqa: E402

ERROS = ('#REF!', '#NAME?', '#DIV/0!', '#VALUE!', '#N/A', '#NUM!', '#NULL!')
XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"
 xmlns:xs="http://www.w3.org/2001/XMLSchema"
 xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def recalc_to(src, outdir, timeout=180):
    """Recalcula `src` e grava o resultado em `outdir`. Devolve o Path do resultado."""
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError('LibreOffice não encontrado (instale em https://www.libreoffice.org/download/)')
    perfil = Path(tempfile.mkdtemp(prefix='lo_perfil_'))
    try:
        (perfil / 'user').mkdir(parents=True)
        (perfil / 'user' / 'registrymodifications.xcu').write_text(XCU, encoding='utf-8')
        cmd = [soffice, f'-env:UserInstallation={perfil.as_uri()}', '--headless',
               '--norestore', '--convert-to', 'xlsx:Calc MS Excel 2007 XML',
               '--outdir', str(outdir), str(src)]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        out = Path(outdir) / (Path(src).stem + '.xlsx')
        if r.returncode != 0 or not out.exists():
            raise RuntimeError(f'LibreOffice falhou: {r.stderr.strip() or r.stdout.strip()}')
        return out
    finally:
        shutil.rmtree(perfil, ignore_errors=True)


def scan_errors(path):
    import openpyxl
    wb_v = openpyxl.load_workbook(path, data_only=True)
    wb_f = openpyxl.load_workbook(path)
    por_tipo = {e: [] for e in ERROS}
    formulas = 0
    for ws in wb_f.worksheets:
        ws_v = wb_v[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    formulas += 1
                v = ws_v[c.coordinate].value
                if isinstance(v, str) and v.strip() in por_tipo:
                    por_tipo[v.strip()].append(f'{ws.title}!{c.coordinate}')
    total = sum(len(v) for v in por_tipo.values())
    return {
        'total_errors': total,
        'total_formulas': formulas,
        'errors': {k: {'count': len(v), 'cells': v[:20]} for k, v in por_tipo.items() if v},
    }


def recalc(path, save=True, timeout=180):
    path = Path(path).resolve()
    tmp = Path(tempfile.mkdtemp(prefix='recalc_'))
    try:
        out = recalc_to(path, tmp, timeout)
        res = scan_errors(out)
        if save:
            shutil.copyfile(out, path)
        res['arquivo'] = str(path)
        res['salvo'] = save
        return res
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def main():
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding='utf-8')
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('file')
    ap.add_argument('--no-save', action='store_true')
    ap.add_argument('--timeout', type=int, default=180)
    a = ap.parse_args()
    try:
        res = recalc(a.file, save=not a.no_save, timeout=a.timeout)
    except Exception as e:  # noqa: BLE001
        print(json.dumps({'status': 'falha', 'erro': str(e)}, ensure_ascii=False))
        return 2
    print(json.dumps(res, ensure_ascii=False, indent=2))
    return 0 if res['total_errors'] == 0 else 1


if __name__ == '__main__':
    sys.exit(main())
