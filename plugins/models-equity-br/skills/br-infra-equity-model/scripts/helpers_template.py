# -*- coding: utf-8 -*-
"""Helpers padronizados para modelos br-infra (openpyxl).
Uso: exec(open('helpers_template.py').read()) no topo de cada build_partN.py,
ajustando Y0 (primeiro ano de projeção), YN (último) e COL0 (coluna de Y0).
Cores: azul=input | preto=fórmula | verde=link entre abas. Comentário com fonte
em todo input. Grade anual em R$ mn."""
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

Y0, YN = 2026, 2056          # ajustar por modelo
COL0 = 4                     # coluna D = Y0 (ajustar se a aba tiver históricos antes)
YEARS = list(range(Y0, YN + 1))

AZUL  = Font(name='Arial', size=10, color='0000FF')
PRETO = Font(name='Arial', size=10, color='000000')
VERDE = Font(name='Arial', size=10, color='008000')
BOLD  = Font(name='Arial', size=10, bold=True)
TIT   = Font(name='Arial', size=12, bold=True)

FILL_SEC  = PatternFill('solid', fgColor='D9E1F2')   # cabeçalho de seção
FILL_AMAR = PatternFill('solid', fgColor='FFF2CC')   # células-check

FMT_MN   = '#,##0;(#,##0)'
FMT_MN1  = '#,##0.0;(#,##0.0)'
FMT_PCT  = '0.0%'
FMT_PRC  = '#,##0.0'
FMT_DATE = 'dd/mm/yyyy'

def col(y):
    """Letra da coluna do ano y na grade padrão (COL0 = Y0)."""
    return get_column_letter(COL0 + (y - Y0))

def w(ws, cell, value, font=PRETO, fmt=None, comment=None, fill=None):
    """Escreve célula com estilo; comentário sempre com fonte/derivação."""
    c = ws[cell]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if comment:
        c.comment = Comment(comment, 'br-infra-skill')
    if fill:
        c.fill = fill
    return c

def sec(ws, row, title):
    """Cabeçalho de seção com fundo."""
    c = ws.cell(row=row, column=1, value=title)
    c.font = BOLD
    c.fill = FILL_SEC
    return c

# Armadilhas conhecidas (não remover estes lembretes):
# - XIRR/XNPV SEM prefixo _xlfn (quebra no recalc LibreOffice).
# - NUNCA insert_cols/insert_rows em grade viva (fórmulas não são reescritas):
#   deletar e recriar a aba, reapontando referências externas.
# - Alinhamento: criar Alignment novo; copiar .alignment de outra célula levanta
#   TypeError (StyleProxy não-hashable).
# - create_file falha se o arquivo existe: remover antes ou usar cat/openpyxl.
# - Fisher: (1+nominal)/(1+ipca)-1. Nunca subtração.
