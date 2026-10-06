#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gate de qualidade antes de QUALQUER entrega.
Uso:
  python3 verify_model.py modelo.xlsx
  python3 verify_model.py modelo.xlsx --tir "Valuation!B57" --ke 0.0769 --upside "Valuation!B43"
  python3 verify_model.py modelo.xlsx --fy "Trimestral!B9:E9=Historico!E12"

Passos: (1) recalcula no LibreOffice via scripts/recalc.py (grava os valores no próprio
arquivo, pois o openpyxl não grava cache) e varre #REF!/#NAME?/#DIV/0! etc.;
(2) exige ZERO erros de fórmula;
(3) CHECK DE ITENS NÃO-CAIXA/NÃO RECORRENTES no EBITDA: varre rótulos (equivalência,
    MtM/valor justo, impairment, baixas, atualizações de ativos) e IMPRIME onde
    aparecem, exigindo confirmação de que o recorrente os expurga (quando
    identificáveis), de que são memos fora do fluxo e de que nada conta duas vezes;
(4) checks de coerência opcionais:
    - TIR real vs Ke quando há upside — coerência TIR x SOTP. ATENÇÃO: o check por
      SINAL falha estruturalmente quando |upside| < ~5% e a alavancagem é alta (o FCFE
      captura o valor do spread da dívida que o SOTP a Ke puro não dá). Nesses casos o
      script só ALERTA (não reprova) e sugere comparar com Ke blended;
    - somas trimestrais = FY (tolerância 0,15 pela derivação por diferença);
(5) lista células-check amarelas (FFF2CC) e seus valores para inspeção final.
ATENÇÃO: nunca rodar o recalc em arquivos com histórico de corrupção (ex.: caso ISAE4)."""
import sys, argparse, json, os, re

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from recalc import recalc  # noqa: E402

NC_PAT = re.compile(r'(equival[êe]ncia patrimonial|valor justo|marca[çc][ãa]o a mercado|\bMtM\b|'
                    r'impairment|atualiza[çc][ãa]o do ativo|n[ãa]o recorrente|baixa de ativ)', re.I)

def parse_ref(ref):
    sheet, cell = ref.split('!')
    return sheet, cell

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('file')
    ap.add_argument('--tir', help='célula da TIR real, ex.: Valuation!B57')
    ap.add_argument('--ke', type=float, help='Ke real (Fisher) para o check de coerência')
    ap.add_argument('--upside', help='célula do upside, ex.: Valuation!B43')
    ap.add_argument('--fy', action='append', default=[],
                    help='check soma=FY: "Aba!B9:E9=Aba2!E12" (repetível)')
    ap.add_argument('--skip-recalc', action='store_true')
    a = ap.parse_args()

    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding='utf-8')
        except (AttributeError, ValueError):
            pass

    ok = True
    if not a.skip_recalc:
        try:
            res = recalc(a.file)
        except Exception as e:  # noqa: BLE001
            print(f'FALHOU: recálculo no LibreOffice — {e}')
            print('Sem recálculo os checks abaixo não enxergam valores. NÃO ENTREGAR.')
            sys.exit(2)
        print(json.dumps(res, ensure_ascii=False, indent=2))
        if res['total_errors']:
            print('FALHOU: erros de fórmula — NÃO ENTREGAR.')
            ok = False

    import openpyxl
    wb = openpyxl.load_workbook(a.file, data_only=True)
    wbf = openpyxl.load_workbook(a.file)

    # ---- (3) CHECK de itens nao-caixa/nao recorrentes no EBITDA ----
    hits = []
    for ws in wbf.worksheets:
        for row in ws.iter_rows(min_col=1, max_col=2):
            for cell in row:
                if isinstance(cell.value, str) and NC_PAT.search(cell.value):
                    hits.append(f'{ws.title}!{cell.coordinate}: "{cell.value.strip()[:70]}"')
    if hits:
        print('\n=== CHECK ITENS NAO-CAIXA/NAO RECORRENTES no EBITDA (responder antes de entregar) ===')
        for h in hits:
            print('  linha nao-caixa/nao recorrente ->', h)
        print('  CONFIRME: (a) calibração e EBITDA RECORRENTE (hist E projetado) expurgam esses itens quando identificáveis;')
        print('            (b) itens não-caixa são memo (fora de FCFF/FCFE/dividendos);')
        print('            (c) nenhum item aparece duas vezes (no fluxo E no bridge);')
        print('            (d) ND/EBITDA e política de dividendos usam o EBITDA RECORRENTE.')
        print('  O EBITDA reportado pode (e deve) aparecer como linha de REFERÊNCIA — nunca como base.')

    # ---- (4) coerência TIR x SOTP ----
    if a.tir and a.ke is not None and a.upside:
        s, c = parse_ref(a.tir); tir = wb[s][c].value
        s, c = parse_ref(a.upside); up = wb[s][c].value
        if isinstance(tir, (int, float)) and isinstance(up, (int, float)):
            coerente = (up < 0) == (tir < a.ke)
            print(f'\nCoerência TIR x SOTP: TIR real {tir:.2%} | Ke {a.ke:.2%} | '
                  f'upside {up:+.1%} -> {"OK" if coerente else "INCOERENTE"}')
            if not coerente:
                if abs(up) < 0.05:
                    print('  AVISO (não reprova): |upside| < 5% — o check por sinal falha '
                          'estruturalmente com alavancagem alta (FCFE captura o spread da '
                          'dívida; SOTP desconta FCFF a Ke puro). Compare a TIR com o Ke '
                          'BLENDED (Ke ajustado pelo mix dívida/equity) e DOCUMENTE no LeiaMe.')
                else:
                    print('  -> o fluxo da TIR provavelmente não deduz obrigações do bridge.')
                    ok = False

    for chk in a.fy:
        left, right = chk.split('=')
        s, rng = left.split('!')
        vals = [c.value for row in wb[s][rng] for c in row]
        soma = sum(v for v in vals if isinstance(v, (int, float)))
        s2, c2 = parse_ref(right); fy = wb[s2][c2].value
        d = soma - (fy or 0)
        print(f'{chk}: soma {soma:,.1f} vs FY {fy:,.1f} | diff {d:+,.2f} '
              f'{"OK" if abs(d) <= 0.15 else "FALHOU"}')
        if abs(d) > 0.15:
            ok = False

    print('\nCélulas-check (fundo amarelo):')
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                f = cell.fill
                if f and f.fgColor and str(f.fgColor.rgb).endswith('FFF2CC'):
                    v = wb[ws.title][cell.coordinate].value
                    print(f'  {ws.title}!{cell.coordinate} = {v}')

    print('\nRESULTADO:', 'APROVADO' if ok else 'REPROVADO — corrigir antes de entregar')
    sys.exit(0 if ok else 1)

if __name__ == '__main__':
    main()
