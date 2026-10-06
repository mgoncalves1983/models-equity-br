#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Teste rápido dos scripts das skills, sem tocar na configuração real do usuário.

Cobre: manifests JSON, compilação dos scripts, ciclo da configuração (primeira rodada
-> set -> status -> get) e, se o LibreOffice estiver instalado, recálculo e
verify_model num modelo mínimo (um limpo e um com #DIV/0!).

Uso: python3 tools/smoke_test.py
"""
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
PLUGIN = RAIZ / 'plugins' / 'models-equity-br'
SKILLS = sorted(p for p in (PLUGIN / 'skills').iterdir() if p.is_dir())
falhas = []


def check(cond, msg):
    print(('  ok    ' if cond else '  FALHA ') + msg)
    if not cond:
        falhas.append(msg)


def run(args, env):
    return subprocess.run([sys.executable] + [str(a) for a in args], env=env,
                          capture_output=True, text=True, encoding='utf-8')


def main():
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding='utf-8')
        except (AttributeError, ValueError):
            pass

    print('manifests')
    for f in (RAIZ / '.claude-plugin' / 'marketplace.json', PLUGIN / '.claude-plugin' / 'plugin.json'):
        try:
            json.loads(f.read_text(encoding='utf-8'))
            check(True, f.relative_to(RAIZ).as_posix())
        except ValueError as e:
            check(False, f'{f.relative_to(RAIZ)}: {e}')

    print('compilação dos scripts')
    for f in sorted(RAIZ.rglob('*.py')):
        if '.git' in f.parts:
            continue
        try:
            compile(f.read_text(encoding='utf-8'), str(f), 'exec')
            check(True, f.relative_to(RAIZ).as_posix())
        except SyntaxError as e:
            check(False, f'{f.relative_to(RAIZ)}: {e}')

    tmp = Path(tempfile.mkdtemp(prefix='smoke_'))
    env = dict(os.environ, MODELS_EQUITY_BR_CONFIG=str(tmp / 'config.json'),
               PYTHONIOENCODING='utf-8')
    ri, saida = tmp / 'ri', tmp / 'modelos'
    ri.mkdir()

    for skill in SKILLS:
        cfg = skill / 'scripts' / 'config.py'
        print(f'configuração ({skill.name})')
        (tmp / 'config.json').unlink(missing_ok=True)
        r = run([cfg, 'status'], env)
        check(r.returncode == 3 and '"configurado": false' in r.stdout,
              'primeira rodada: status devolve "não configurado" (código 3)')
        r = run([cfg, 'set', '--pasta-ri', tmp / 'nao_existe', '--pasta-saida', saida], env)
        check(r.returncode == 1, 'set recusa pasta de RI inexistente sem --criar')
        r = run([cfg, 'set', '--pasta-ri', ri, '--pasta-fundamentos', '', '--pasta-saida', saida], env)
        check(r.returncode == 0 and saida.is_dir(), 'set grava a configuração e cria a pasta de saída')
        r = run([cfg, 'status'], env)
        check(r.returncode == 0 and '"configurado": true' in r.stdout, 'status devolve "configurado"')
        r = run([cfg, 'get', 'pasta_ri'], env)
        check(r.returncode == 0 and Path(r.stdout.strip()) == ri.resolve(), 'get pasta_ri')
        r = run([cfg, 'deps'], env)
        print('        ' + r.stdout.strip().replace('\n', '\n        '))

    sys.path.insert(0, str(SKILLS[0] / 'scripts'))
    from config import find_soffice  # noqa: E402
    if not find_soffice():
        print('recálculo: PULADO (LibreOffice não instalado)')
    else:
        import openpyxl
        print('recálculo e verify_model')
        for nome, div0 in (('limpo', False), ('com_erro', True)):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'Valuation'
            ws['A1'], ws['B1'] = 'EBITDA', 100
            ws['A2'], ws['B2'] = 'Margem', '=B1/0' if div0 else '=B1/200'
            wb.save(saida / f'{nome}.xlsx')
        for skill in SKILLS:
            sc = skill / 'scripts'
            r = run([sc / 'recalc.py', saida / 'limpo.xlsx', '--no-save'], env)
            check(r.returncode == 0 and '"total_errors": 0' in r.stdout,
                  f'{skill.name}: recalc.py sem erros no modelo limpo')
            r = run([sc / 'recalc.py', saida / 'com_erro.xlsx', '--no-save'], env)
            check(r.returncode == 1 and '#DIV/0!' in r.stdout,
                  f'{skill.name}: recalc.py acusa #DIV/0!')
            r = run([sc / 'verify_model.py', saida / 'limpo.xlsx'], env)
            check(r.returncode == 0 and 'APROVADO' in r.stdout,
                  f'{skill.name}: verify_model.py aprova o modelo limpo')
            r = run([sc / 'verify_model.py', saida / 'com_erro.xlsx'], env)
            check(r.returncode == 1 and 'REPROVADO' in r.stdout,
                  f'{skill.name}: verify_model.py reprova o modelo com erro')

    print()
    if falhas:
        print(f'{len(falhas)} falha(s).')
        return 1
    print('Tudo certo.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
